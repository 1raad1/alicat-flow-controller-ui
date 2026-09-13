import csv
import os
from pathlib import Path
import tempfile
import threading
import time
import unittest
from datetime import datetime, timedelta
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QApplication

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from flow_controller.core.graph_history import GraphHistory, HistoryExportSnapshot
from flow_controller.core.history_export import HistoryExporter


def sample(value):
    return {
        "flow": value,
        "sp": value + 1,
        "press": value + 2,
        "temp": value + 3,
        "internal_error": value + 4,
        "valve_drives": (value + 5,),
    }


class HistorySnapshotTests(unittest.TestCase):
    def test_snapshot_is_immutable_and_detached_from_live_history(self):
        history = GraphHistory(limit=3)
        history.set_units([1])
        started = datetime(2026, 1, 1)
        history.push(1, started, {1: sample(10)})
        snapshot = history.export_snapshot(metric_keys=["flow"])

        history.push(2, started + timedelta(seconds=1), {1: sample(20)})
        history.clear(generation=2)

        self.assertIsInstance(snapshot, HistoryExportSnapshot)
        self.assertEqual(snapshot.header, ("time_s", "U1_flow_SLPM"))
        self.assertEqual(tuple(snapshot.iter_rows()), ((0.0, 10),))
        with self.assertRaises(AttributeError):
            snapshot.times = ()


class HistoryExporterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def snapshot(self):
        return HistoryExportSnapshot(
            header=("time_s", "U1_flow_SLPM", "U2_flow_SLPM"),
            times=(0.0, 1.0, 2.0),
            columns=((1.25, None, 3.5), (9.0,)),
        )

    def wait_until(self, predicate, timeout=3.0):
        deadline = time.monotonic() + timeout
        while not predicate() and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(0.001)
        self.app.processEvents()
        self.assertTrue(predicate(), "timed out waiting for background export")

    def test_csv_completion_writes_headers_blanks_and_tail_alignment(self):
        exporter = HistoryExporter()
        self.addCleanup(exporter.shutdown)
        states = []
        completed = []
        exporter.active_changed.connect(states.append)
        exporter.completed.connect(lambda path, count: completed.append((path, count)))
        destination = self.root / "history.csv"

        self.assertTrue(exporter.start(self.snapshot(), destination))
        self.wait_until(lambda: bool(completed))

        with destination.open(newline="", encoding="utf-8") as handle:
            rows = list(csv.reader(handle))
        self.assertEqual(rows, [
            ["time_s", "U1_flow_SLPM", "U2_flow_SLPM"],
            ["0.0", "1.25", ""],
            ["1.0", "", ""],
            ["2.0", "3.5", "9.0"],
        ])
        self.assertEqual(completed, [(str(destination), 3)])
        self.assertEqual(states, [True, False])
        self.assertFalse(exporter.active)
        self.assertTrue(exporter.shutdown())

    @unittest.skipIf(load_workbook is None, "openpyxl is not installed")
    def test_xlsx_completion_uses_live_data_sheet_and_blank_cells(self):
        exporter = HistoryExporter()
        self.addCleanup(exporter.shutdown)
        completed = []
        exporter.completed.connect(lambda path, count: completed.append((path, count)))
        destination = self.root / "history.xlsx"

        self.assertTrue(exporter.start(self.snapshot(), destination))
        self.wait_until(lambda: bool(completed))

        workbook = load_workbook(destination, read_only=True, data_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Live Data"])
        rows = list(workbook["Live Data"].iter_rows(
            min_col=1, max_col=3, values_only=True))
        self.assertEqual(rows, [
            ("time_s", "U1_flow_SLPM", "U2_flow_SLPM"),
            (0, 1.25, None),
            (1, None, None),
            (2, 3.5, 9),
        ])
        self.assertEqual(completed[0][1], 3)

    def test_writer_error_preserves_existing_target_and_removes_temporary(self):
        destination = self.root / "history.csv"
        destination.write_text("original", encoding="utf-8")
        exporter = HistoryExporter()
        self.addCleanup(exporter.shutdown)
        errors = []
        exporter.failed.connect(errors.append)

        with patch("flow_controller.core.history_export._write_csv",
                   side_effect=OSError("disk full")):
            self.assertTrue(exporter.start(self.snapshot(), destination))
            self.wait_until(lambda: bool(errors))

        self.assertIn("disk full", errors[0])
        self.assertEqual(destination.read_text(encoding="utf-8"), "original")
        self.assertEqual(list(self.root.glob("history.csv.*.tmp")), [])

    def test_blocked_export_is_singleton_cancellable_and_does_not_block_events(self):
        destination = self.root / "history.csv"
        destination.write_text("original", encoding="utf-8")
        entered = threading.Event()
        release = threading.Event()
        ui_progressed = []
        cancelled = []
        exporter = HistoryExporter()

        def cleanup():
            release.set()
            self.wait_until(lambda: not exporter.active)
            exporter.shutdown()

        self.addCleanup(cleanup)

        def blocked_writer(_snapshot, path, _cancelled):
            Path(path).write_text("partial", encoding="utf-8")
            entered.set()
            release.wait(3)
            return 0

        exporter.cancelled.connect(lambda: cancelled.append(True))
        with patch("flow_controller.core.history_export._write_csv",
                   side_effect=blocked_writer):
            self.assertTrue(exporter.start(self.snapshot(), destination))
            self.wait_until(entered.is_set)
            self.assertFalse(exporter.start(self.snapshot(), self.root / "second.csv"))
            QTimer.singleShot(0, lambda: ui_progressed.append(True))
            self.wait_until(lambda: bool(ui_progressed))

            self.assertFalse(exporter.shutdown())
            self.assertTrue(exporter.active)
            self.assertEqual(destination.read_text(encoding="utf-8"), "original")
            release.set()
            self.wait_until(lambda: bool(cancelled))

        self.assertFalse(exporter.active)
        self.assertEqual(destination.read_text(encoding="utf-8"), "original")
        self.assertEqual(list(self.root.glob("history.csv.*.tmp")), [])
        self.assertTrue(exporter.shutdown())

    @unittest.skipIf(load_workbook is None, "openpyxl is not installed")
    def test_cancelled_xlsx_removes_destination_and_worksheet_temporaries(self):
        from openpyxl.worksheet._writer import ALL_TEMP_FILES

        destination = self.root / "history.xlsx"
        destination.write_bytes(b"original")
        entered = threading.Event()
        release = threading.Event()
        cancelled = []
        exporter = HistoryExporter()
        before = set(ALL_TEMP_FILES)

        class BlockingSnapshot:
            header = ("time_s", "flow")

            def iter_rows(self):
                entered.set()
                release.wait(3)
                yield (0.0, 1.0)

        def cleanup():
            release.set()
            self.wait_until(lambda: not exporter.active)
            exporter.shutdown()

        self.addCleanup(cleanup)
        exporter.cancelled.connect(lambda: cancelled.append(True))
        self.assertTrue(exporter.start(BlockingSnapshot(), destination))
        self.wait_until(entered.is_set)
        self.assertFalse(exporter.shutdown())
        release.set()
        self.wait_until(lambda: bool(cancelled))

        self.assertEqual(destination.read_bytes(), b"original")
        self.assertEqual(set(ALL_TEMP_FILES), before)
        self.assertEqual(list(self.root.glob("history.xlsx.*.tmp")), [])

    @unittest.skipIf(load_workbook is None, "openpyxl is not installed")
    def test_xlsx_row_error_preserves_target_and_cleans_temporaries(self):
        from openpyxl.worksheet._writer import ALL_TEMP_FILES

        destination = self.root / "history.xlsx"
        destination.write_bytes(b"original")
        errors = []
        exporter = HistoryExporter()
        self.addCleanup(exporter.shutdown)
        before = set(ALL_TEMP_FILES)

        class FailingSnapshot:
            header = ("time_s", "flow")

            def iter_rows(self):
                yield (0.0, 1.0)
                raise OSError("source failed")

        exporter.failed.connect(errors.append)
        self.assertTrue(exporter.start(FailingSnapshot(), destination))
        self.wait_until(lambda: bool(errors))

        self.assertIn("source failed", errors[0])
        self.assertEqual(destination.read_bytes(), b"original")
        self.assertEqual(set(ALL_TEMP_FILES), before)
        self.assertEqual(list(self.root.glob("history.xlsx.*.tmp")), [])


if __name__ == "__main__":
    unittest.main()

"""Atomic background writers for immutable graph-history snapshots."""

from __future__ import annotations

import csv
import os
from pathlib import Path
import tempfile

from PySide6.QtCore import QObject, QThread, Signal

from .graph_history import HistoryExportSnapshot


class _ExportCancelled(Exception):
    pass


def _check_cancelled(cancelled):
    if cancelled():
        raise _ExportCancelled


def _write_csv(snapshot, path, cancelled):
    count = 0
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(snapshot.header)
        for row in snapshot.iter_rows():
            _check_cancelled(cancelled)
            writer.writerow(row)
            count += 1
    return count


def _write_xlsx(snapshot, path, cancelled):
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("Excel export requires openpyxl.") from exc

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Live Data")
    count = 0
    saved = False
    try:
        sheet.append(snapshot.header)
        for row in snapshot.iter_rows():
            _check_cancelled(cancelled)
            sheet.append(row)
            count += 1
        _check_cancelled(cancelled)
        workbook.save(path)
        saved = True
    finally:
        if not saved:
            # Workbook.close() does not close an unsaved write-only worksheet
            # or remove its temporary XML file in openpyxl 3.1. Finish the
            # sheet writer and discard that file explicitly on cancellation or
            # failure; the public API has no unsaved-workbook cleanup method.
            writer = getattr(sheet, "_writer", None)
            if writer is not None:
                try:
                    sheet.close()
                except Exception:
                    try:
                        writer.close()
                    except Exception:
                        pass
                try:
                    if os.path.exists(writer.out):
                        writer.cleanup()
                except (OSError, ValueError):
                    pass
        workbook.close()
    return count


def _write_atomic(snapshot, path, cancelled):
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(
        prefix=destination.name + ".", suffix=".tmp", dir=destination.parent)
    os.close(descriptor)
    try:
        _check_cancelled(cancelled)
        suffix = destination.suffix.lower()
        if suffix == ".csv":
            count = _write_csv(snapshot, temporary, cancelled)
        elif suffix == ".xlsx":
            count = _write_xlsx(snapshot, temporary, cancelled)
        else:
            raise ValueError("History export path must end in .csv or .xlsx.")
        _check_cancelled(cancelled)
        os.replace(temporary, destination)
        return str(destination), count
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


class _HistoryExportWorker(QThread):
    def __init__(self, snapshot, path, parent=None):
        super().__init__(parent)
        self.snapshot = snapshot
        self.path = str(path)
        self.outcome = None

    def run(self):
        try:
            path, count = _write_atomic(
                self.snapshot, self.path, self.isInterruptionRequested)
            self.outcome = ("completed", path, count)
        except _ExportCancelled:
            self.outcome = ("cancelled",)
        except Exception as exc:
            self.outcome = ("failed", str(exc))


class HistoryExporter(QObject):
    """Own one bounded export job and report its result on the owner thread."""

    active_changed = Signal(bool)
    completed = Signal(str, int)
    failed = Signal(str)
    cancelled = Signal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self._worker = None

    @property
    def active(self):
        return self._worker is not None

    def start(self, snapshot: HistoryExportSnapshot, path):
        if self.active:
            return False
        worker = _HistoryExportWorker(snapshot, path, self)
        self._worker = worker
        worker.finished.connect(self._worker_finished)
        self.active_changed.emit(True)
        worker.start()
        return True

    def cancel(self):
        if self._worker is None:
            return False
        self._worker.requestInterruption()
        return True

    def _worker_finished(self):
        worker = self._worker
        if worker is None or worker.isRunning():
            return
        self._worker = None
        outcome = worker.outcome or ("failed", "History export ended without a result.")
        worker.deleteLater()
        self.active_changed.emit(False)
        if outcome[0] == "completed":
            self.completed.emit(outcome[1], outcome[2])
        elif outcome[0] == "cancelled":
            self.cancelled.emit()
        else:
            self.failed.emit(outcome[1])

    def shutdown(self):
        worker = self._worker
        if worker is None:
            return True
        worker.requestInterruption()
        if not worker.wait(100):
            return False
        self._worker_finished()
        return True
